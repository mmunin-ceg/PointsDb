import psycopg2
from psycopg2.extras import execute_values
from openpyxl import load_workbook
import openpyxl

# === Configuration ===
EXCEL_PATH = "PPC Solar Master Points List.xlsx"
PG_CONN_INFO = {
    "dbname": "points",
    "user": "mmunin",
    "password": "Password1!",
    "host": "localhost",
    "port": 5432,
}
DEFAULT_PROVIDER = "CEG"
DEFAULT_PROTOCOL = "Modbus"
DEFAULT_VERSION = "unversioned"  # fallback

# Expected column mapping (case-insensitive match)
EXPECTED_COLUMNS = [
    "List Index", "Availability", "Object", "Point Description", "Data Type",
    "DNP index", "Bit", "Alarm State", "On (1) State", "Off (0) State", "Units",
    "Scalar", "Analog Alarm Limits", "Alarm Profile", "Historian Profile",
    "HMI Category", "RTAC Point Name", "Enumeration Table", "Comments"
]

# === Helper: Normalize Excel column names ===
def normalize(header):
    if not header:
        return ""
    return str(header).lower().strip()

# === PostgreSQL connection ===
conn = psycopg2.connect(**PG_CONN_INFO)
cur = conn.cursor()

# Ensure provider exists
cur.execute("""
    INSERT INTO api_provider (name, is_internal)
    VALUES (%s, true)
    ON CONFLICT (name) DO NOTHING
""", (DEFAULT_PROVIDER,))
conn.commit()

# === Load Excel workbook ===
wb = load_workbook(EXCEL_PATH, data_only=True, read_only=False)
skip = False
for sheet in wb.worksheets:
    if sheet.title == "PPC to Ignition - SG3600":
        skip = False
    if skip:
        continue
        
    print(f"Processing sheet: {sheet.title}")
    if sheet.title == "Data Flows":
        print("  Skipped: 'Data Flows' sheet.")
        continue
    if not sheet._tables:
        print(f"  Skipped: no Excel tables found on this sheet.")
        continue
    
    table = next(iter(sheet.tables.values()))  # Only consider the first table per sheet
    ref = table.ref
    min_col, min_row, max_col, max_row = openpyxl.utils.range_boundaries(ref)
    print(f"  Table range: {min_row}-{max_row}, {min_col}-{max_col}")
    headers = [sheet.cell(row=min_row, column=c).value for c in range(min_col, max_col + 1)]
    headers_norm = [normalize(h) if h else "" for h in headers]
    col_map = {normalize(h): i+1 for i, h in enumerate(headers)}

    provider = input("Enter the provider name: ").strip()
    protocol = input("Enter the protocol name: ").strip()
    interface_name = sheet.title.strip()
    version = DEFAULT_VERSION

    # If provider does not exist, insert it
    cur.execute("""
        INSERT INTO api_provider (name, is_internal)
        VALUES (%s, true)
        ON CONFLICT (name) DO NOTHING
        RETURNING id
    """, (provider,))
    if cur.rowcount == 0:
        cur.execute("SELECT id FROM api_provider WHERE name = %s", (provider,))
    provider_id = cur.fetchone()[0]
    conn.commit()

    # If interface does not exist, insert it
    cur.execute("""
        INSERT INTO data_interface (name, protocol, provider_id)
        VALUES (%s, %s, %s)
        ON CONFLICT (name, protocol, provider_id) DO NOTHING
        RETURNING id
    """, (interface_name, protocol, provider_id))
    if cur.rowcount == 0:
        cur.execute("""
            SELECT id FROM data_interface 
            WHERE name = %s AND protocol = %s AND provider_id = %s
        """, (interface_name, protocol, provider_id))
    interface_id = cur.fetchone()[0]
    conn.commit()

    # Get or create map version
    cur.execute("""
        INSERT INTO map_version (version, interface_id)
        VALUES (%s, %s)
        ON CONFLICT (version, interface_id) DO NOTHING
        RETURNING id
    """, (version, interface_id))
    if cur.rowcount == 0:
        cur.execute("""
            SELECT id FROM map_version 
            WHERE version = %s AND interface_id = %s
        """, (version, interface_id))
    version_id = cur.fetchone()[0]
    conn.commit()

    # Process each row in the table
    for r in range(min_row + 1, max_row + 1):
        if sheet.cell(row=r, column=1).value is None:
            continue
            
        point_name = sheet.cell(row=r, column=col_map['rtac point name']).value
        object_name = sheet.cell(row=r, column=col_map['object']).value
        
        if not point_name:
            print(f"  Skipping row {r}: No point name")
            continue
            
        # Check if "register" or "dnp index" exists and use accordingly
        if "register" in col_map:
            register = sheet.cell(row=r, column=col_map['register']).value
        elif "dnp index" in col_map:
            register = sheet.cell(row=r, column=col_map['dnp index']).value
        else:
            register = None        
        data_type = sheet.cell(row=r, column=col_map['data type']).value
        if "bit" in col_map:
            bit_offset = sheet.cell(row=r, column=col_map['bit']).value
        elif "modbus bit" in col_map:
            bit_offset = sheet.cell(row=r, column=col_map['modbus bit']).value
        
        units = sheet.cell(row=r, column=col_map['units']).value 
        scale = sheet.cell(row=r, column=col_map['scalar']).value
        alarm_state = sheet.cell(row=r, column=col_map['alarm state']).value
        on_state = sheet.cell(row=r, column=col_map['on (1) state']).value
        off_state = sheet.cell(row=r, column=col_map['off (0) state']).value
        alarm_limits = sheet.cell(row=r, column=col_map['analog alarm limits']).value
        alarm_profile = sheet.cell(row=r, column=col_map['alarm profile']).value
        enumeration_table = sheet.cell(row=r, column=col_map['enumeration table']).value
        comments = sheet.cell(row=r, column=col_map['comments']).value

        # Insert point
        cur.execute("""
            INSERT INTO map_point (
                version_id, point_name, object_name, register, data_type,
                bit_offset, units, scale, alarm_state, on_state,
                off_state, alarm_limits, alarm_profile, enumeration_table, comments
            )
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            ON CONFLICT (point_name, version_id) DO UPDATE SET
                object_name = EXCLUDED.object_name,
                register = EXCLUDED.register,
                data_type = EXCLUDED.data_type,
                bit_offset = EXCLUDED.bit_offset,
                units = EXCLUDED.units,
                scale = EXCLUDED.scale,
                alarm_state = EXCLUDED.alarm_state,
                on_state = EXCLUDED.on_state,
                off_state = EXCLUDED.off_state,
                alarm_limits = EXCLUDED.alarm_limits,
                alarm_profile = EXCLUDED.alarm_profile,
                enumeration_table = EXCLUDED.enumeration_table,
                comments = EXCLUDED.comments
        """, (
            version_id, point_name, object_name, register, data_type,
            bit_offset, units, scale, alarm_state, on_state,
            off_state, alarm_limits, alarm_profile, enumeration_table, comments
        ))

    conn.commit()
    print(f"  Processed sheet: {sheet.title}")

conn.close()
print("Done!")
